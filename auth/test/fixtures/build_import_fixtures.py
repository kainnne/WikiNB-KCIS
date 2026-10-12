from pathlib import Path
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


def add_hyperlink(paragraph, text, url):
    relation_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1B4F9C")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([properties, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def main():
    output = Path(sys.argv[1]).resolve()
    output.mkdir(parents=True, exist_ok=True)
    image_path = output / "fixture-image.png"

    image = Image.new("RGB", (900, 360), "#edf4ff")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((25, 25, 875, 335), radius=35, outline="#6b3d9a", width=8)
    draw.text((70, 145), "KCIS IMPORT FIXTURE", fill="#1b4f9c")
    image.save(image_path)

    document = Document()
    document.add_heading("DOCX Import Fixture", level=1)
    document.add_paragraph("This paragraph verifies semantic text conversion.")
    link_paragraph = document.add_paragraph("Reference: ")
    add_hyperlink(link_paragraph, "KCIS example link", "https://example.com/kcis")
    document.add_picture(str(image_path), width=Inches(5.8))
    document.add_paragraph("First item", style="List Bullet")
    document.add_paragraph("Second item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Criterion"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "Images and links"
    table.cell(1, 1).text = "Preserved"
    document.save(output / "import-fixture.docx")

    pdf_path = output / "import-fixture.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4)
    width, height = A4
    pdf.setTitle("PDF Import Fixture")
    pdf.setFont("Helvetica-Bold", 22)
    pdf.setFillColorRGB(0.10, 0.31, 0.61)
    pdf.drawString(60, height - 80, "PDF Import Fixture")
    pdf.setFont("Helvetica", 12)
    pdf.setFillColorRGB(0.15, 0.15, 0.20)
    pdf.drawString(60, height - 115, "Searchable PDF text for WikiNB and Gemini.")
    pdf.setFillColorRGB(0.42, 0.24, 0.60)
    pdf.drawString(60, height - 145, "https://example.com/kcis-pdf")
    pdf.linkURL("https://example.com/kcis-pdf", (60, height - 150, 260, height - 132))
    pdf.drawImage(str(image_path), 60, height - 430, width=470, height=188, preserveAspectRatio=True)
    pdf.showPage()
    pdf.save()


if __name__ == "__main__":
    main()
